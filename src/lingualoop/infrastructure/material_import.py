from __future__ import annotations

import re
import zipfile
from html.parser import HTMLParser
from xml.etree import ElementTree
from posixpath import dirname, join, normpath
from dataclasses import dataclass, field
from pathlib import Path

from lingualoop.core.domain import MaterialSourceType, SubtitleCue


@dataclass(frozen=True)
class ImportedMaterial:
    title: str
    raw_content: str
    normalized_content: str
    source_type: MaterialSourceType
    source_path: str
    diagnostics: list[str]
    subtitle_cues: list[SubtitleCue] = field(default_factory=list)


def import_material_file(path: Path | str) -> ImportedMaterial:
    file_path = Path(path).expanduser()
    if not file_path.is_file():
        raise FileNotFoundError(f"Material file does not exist: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix in {".mp4", ".mkv", ".webm", ".mov", ".avi"}:
        sidecar = next((file_path.with_suffix(ext) for ext in (".srt", ".vtt")
                        if file_path.with_suffix(ext).is_file()), None)
        if sidecar is None:
            raise ValueError("视频导入需要同名 .srt 或 .vtt 字幕；当前只分析字幕文字，尚未识别视频画面。")
        imported = import_material_file(sidecar)
        return ImportedMaterial(title=file_path.stem, raw_content=imported.raw_content,
            normalized_content=imported.normalized_content, source_type=MaterialSourceType.SUBTITLE,
            source_path=str(file_path.resolve()), subtitle_cues=imported.subtitle_cues,
            diagnostics=[*imported.diagnostics, "Imported sidecar subtitles only; video frames were not analyzed."])
    if file_path.stat().st_size > 60 * 1024 * 1024:
        raise ValueError("导入文件超过 60 MiB，请按卷拆分后导入。")
    cues: list[SubtitleCue] = []
    metadata_title = None
    if suffix in {".md", ".markdown", ".txt", ".srt", ".vtt"}:
        data = file_path.read_bytes()
        raw = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
        source_type = MaterialSourceType.ARTICLE
        if suffix in {".srt", ".vtt"}:
            import pysubs2

            subtitles = pysubs2.SSAFile.from_string(raw, format_=suffix[1:])
            cues = [SubtitleCue(start_ms=max(0, event.start), end_ms=max(0, event.end), text=event.plaintext.strip())
                    for event in subtitles if event.plaintext.strip() and not event.is_comment]
            source_type = MaterialSourceType.SUBTITLE
    elif suffix == ".docx":
        raw = _read_docx(file_path)
        source_type = MaterialSourceType.NOTE
    elif suffix == ".epub":
        raw, metadata_title = _read_epub(file_path)
        source_type = MaterialSourceType.ARTICLE
    elif suffix == ".xlsx":
        raw = _read_xlsx(file_path)
        source_type = MaterialSourceType.ARTICLE
    else:
        raise ValueError("Supported materials: .txt, .md, .markdown, .docx, .srt, .vtt, .epub, .xlsx; video with sidecar subtitles.")

    if suffix in {".srt", ".vtt"}:
        normalized = "\n".join(cue.text for index, cue in enumerate(cues)
                               if index == 0 or cue.text != cues[index - 1].text)
        diagnostics = ["Subtitle text normalized; raw source and cue timestamps preserved."]
    else:
        normalized, diagnostics = normalize_material(raw, suffix)
    if not normalized:
        raise ValueError("The selected material does not contain readable text.")
    return ImportedMaterial(
        title=metadata_title or file_path.stem.replace("_", " ").replace("-", " ").strip() or "Imported material",
        raw_content=raw,
        normalized_content=normalized,
        source_type=source_type,
        source_path=str(file_path.resolve()),
        diagnostics=diagnostics,
        subtitle_cues=cues,
    )


def normalize_material(raw: str, suffix: str = ".md") -> tuple[str, list[str]]:
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    diagnostics: list[str] = []
    if suffix in {".md", ".markdown"}:
        text = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", text, flags=re.DOTALL)
        text = re.sub(r"```[\w-]*\n(.*?)```", r"\1", text, flags=re.DOTALL)
        text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
        diagnostics.append("Removed Markdown formatting while preserving readable text.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    normalized = "\n".join(line.strip() for line in text.splitlines()).strip()
    if normalized != raw.strip():
        diagnostics.append("Created a lesson-ready normalized copy; raw content is preserved.")
    return normalized, diagnostics


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX import requires the python-docx dependency.") from exc

    document = Document(path)
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return "\n\n".join(blocks)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    _check_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    blocks = []
    count = 0
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                count += 1
                if count > 100000:
                    raise ValueError("Excel 超过 100000 行，请按工作表或章节拆分。")
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    rows.append(" | ".join(values).rstrip(" |"))
            if rows:
                blocks.append(f"【{sheet.title}】\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def _check_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 120 * 1024 * 1024:
            raise ValueError("解压后的导入材料超过 120 MiB，请拆分文件。")


class _ReadableHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "rt", "rp"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "h1", "h2", "h3", "li"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "rt", "rp"}:
            self.hidden = max(0, self.hidden - 1)
        elif tag in {"p", "div", "h1", "h2", "h3", "li"} and not self.hidden:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _read_epub(path: Path) -> tuple[str, str | None]:
    _check_archive(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("META-INF/container.xml"))
        rootfile = next((element.attrib.get("full-path") for element in root.iter()
                         if element.tag.endswith("rootfile")), None)
        if not rootfile:
            raise ValueError("EPUB 缺少可读取的书目。")
        package = ElementTree.fromstring(archive.read(rootfile))
        title = next((element.text.strip() for element in package.iter()
                      if element.tag == '{http://purl.org/dc/elements/1.1/}title' and element.text and element.text.strip()), None)
        manifest = {element.attrib.get("id"): element.attrib.get("href")
                    for element in package.iter() if element.tag.endswith("}item")}
        spine = [element.attrib.get("idref") for element in package.iter() if element.tag.endswith("itemref")]
        blocks = []
        for idref in spine:
            href = manifest.get(idref)
            if not href:
                continue
            member = normpath(join(dirname(rootfile), href.split("#")[0]))
            parser = _ReadableHTML()
            parser.feed(archive.read(member).decode("utf-8-sig"))
            text = "".join(parser.parts).strip()
            if text:
                blocks.append(text)
        return "\n\n".join(blocks), title
